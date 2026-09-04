import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    COLOR_DARK_BG = RGBColor(15, 23, 42)       # #0F172A
    COLOR_LIGHT_BG = RGBColor(248, 250, 252)   # #F8FAFC
    COLOR_CARD_BG = RGBColor(255, 255, 255)    # #FFFFFF
    COLOR_PRIMARY = RGBColor(13, 148, 136)     # Teal/Emerald #0D9488
    COLOR_PRIMARY_DARK = RGBColor(15, 118, 110)
    COLOR_TEXT_DARK = RGBColor(15, 23, 42)
    COLOR_TEXT_MUTED = RGBColor(71, 85, 105)   # Slate 600
    COLOR_ACCENT = RGBColor(217, 119, 6)       # Amber #D97706
    COLOR_BORDER = RGBColor(226, 232, 240)     # Slate 200

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, subtitle, title):
        # Subtitle / Kicker
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle.upper()
        p_sub.font.size = Pt(11)
        p_sub.font.bold = True
        p_sub.font.color.rgb = COLOR_PRIMARY

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.6))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_DARK

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    def add_image_card(slide, img_path, left, top, width, height, caption_title, caption_desc):
        # Card Background
        add_card(slide, left, top, width, height, bg_color=COLOR_CARD_BG, border_color=COLOR_BORDER)

        # Image placement slightly inside border
        img_top = top + Inches(0.1)
        img_left = left + Inches(0.1)
        img_width = width - Inches(0.2)
        img_height = height - Inches(1.1)

        if os.path.exists(img_path):
            slide.shapes.add_picture(img_path, img_left, img_top, width=img_width, height=img_height)
        else:
            print(f"WARNING: Image not found at {img_path}")

        # Caption Box
        cap_box = slide.shapes.add_textbox(left + Inches(0.2), top + height - Inches(0.95), width - Inches(0.4), Inches(0.8))
        tf = cap_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = caption_title
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_DARK

        p2 = tf.add_paragraph()
        p2.text = caption_desc
        p2.font.size = Pt(10)
        p2.font.color.rgb = COLOR_TEXT_MUTED

    # ==================== SLIDE 1: Title Slide ====================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, COLOR_DARK_BG)

    # Accent bar on left
    accent_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.2), Inches(0.15), Inches(3.2))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = COLOR_PRIMARY
    accent_bar.line.fill.background()

    # Title text box
    tb1 = slide1.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(6.5), Inches(3.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "KASHMIR'S HIDDEN VALHALLA"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    p2 = tf1.add_paragraph()
    p2.text = "Gurez Valley"
    p2.font.size = Pt(44)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(255, 255, 255)
    p2.space_before = Pt(10)

    p3 = tf1.add_paragraph()
    p3.text = "The Unexplored Crown of High Altitude Kashmir"
    p3.font.size = Pt(18)
    p3.font.color.rgb = RGBColor(203, 213, 225) # Slate 300
    p3.space_before = Pt(10)

    p4 = tf1.add_paragraph()
    p4.text = "A comprehensive exploration of geography, the Dard-Shina culture, Habba Khatoon peak, Kishanganga river, and sustainable eco-tourism."
    p4.font.size = Pt(12)
    p4.font.color.rgb = RGBColor(148, 163, 184) # Slate 400
    p4.space_before = Pt(15)

    # Title Hero Image Card
    add_image_card(
        slide1,
        "presentation/images/Habba_Khatoon_Peak.jpg",
        Inches(8.0), Inches(1.2), Inches(4.5), Inches(5.1),
        "Habba Khatoon Mountain",
        "The pyramid peak towering over Dawar"
    )

    # ==================== SLIDE 2: Location & Geography ====================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, COLOR_LIGHT_BG)
    add_header(slide2, "Location & Geography", "Gateway to High Altitude Wonder")

    points_s2 = [
        ("Geographic Position", "Located in Bandipora district, Northern Kashmir at ~8,000 ft (2,400m) altitude surrounded by towering Himalayan peaks."),
        ("Razdan Pass (11,672 ft)", "High mountain pass connecting Srinagar/Bandipora to Gurez. Covered in deep snow during winter, preserving pristine isolation."),
        ("Alpine Ecosystem", "Dense pine forests, sub-alpine meadows, and pristine river valleys hosting rare wildlife species.")
    ]

    for i, (title, desc) in enumerate(points_s2):
        top_pos = Inches(1.6 + i * 1.7)
        add_card(slide2, Inches(0.8), top_pos, Inches(6.5), Inches(1.5))

        tb = slide2.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.15), Inches(6.1), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_DARK

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_before = Pt(4)

    add_image_card(
        slide2,
        "presentation/images/Gurez_Valley_Overview.jpg",
        Inches(7.6), Inches(1.6), Inches(4.9), Inches(4.9),
        "Pristine Alpine Landscape",
        "Dramatic valley views along the Kishanganga river corridor in Bandipora district."
    )

    # ==================== SLIDE 3: Habba Khatoon ====================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, COLOR_LIGHT_BG)
    add_header(slide3, "Natural Wonders", "The Majestic Habba Khatoon Peak")

    add_image_card(
        slide3,
        "presentation/images/Habba_Khatoon_Peak.jpg",
        Inches(0.8), Inches(1.6), Inches(5.2), Inches(4.9),
        "Pyramid Peak of Habba Khatoon",
        "Iconic mountain backdrop overlooking the valley town of Dawar."
    )

    points_s3 = [
        ("Legend of the Poetess Queen", "Named after Habba Khatoon, the 16th-century 'Nightingale of Kashmir'. Legend holds she wandered these slopes mourning her exiled husband, King Yusuf Shah Chak."),
        ("Geological Significance", "A striking pyramid-shaped monolith rising sharply above the Kishanganga riverbed, creating a distinct geographical landmark."),
        ("Golden Hour Spectacle", "At sunset, the limestone peak catches brilliant amber and gold light, making it a world-renowned highlight for photographers and visitors.")
    ]

    for i, (title, desc) in enumerate(points_s3):
        top_pos = Inches(1.6 + i * 1.7)
        add_card(slide3, Inches(6.3), top_pos, Inches(6.2), Inches(1.5))

        tb = slide3.shapes.add_textbox(Inches(6.5), top_pos + Inches(0.15), Inches(5.8), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_DARK

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_before = Pt(4)

    # ==================== SLIDE 4: Kishanganga River & Tulail Valley ====================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, COLOR_LIGHT_BG)
    add_header(slide4, "Hydrology & Valleys", "Kishanganga River & Tulail Valley")

    points_s4 = [
        ("Lifeline of Gurez", "The crystal-clear Kishanganga River (Neelum) flows westward through the valley, nourishing fertile banks and supporting local trout fishing."),
        ("Tulail Sub-Valley", "Extending further east, Tulail represents an even more untouched frontier with traditional log homes and pristine alpine wilderness."),
        ("Hydroelectric & Ecology", "Host to the 330 MW Kishanganga Hydroelectric Project, balancing clean energy production with fragile Himalayan ecology.")
    ]

    for i, (title, desc) in enumerate(points_s4):
        top_pos = Inches(1.6 + i * 1.7)
        add_card(slide4, Inches(0.8), top_pos, Inches(6.5), Inches(1.5))

        tb = slide4.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.15), Inches(6.1), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_DARK

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_before = Pt(4)

    add_image_card(
        slide4,
        "presentation/images/Kishanganga_River.jpg",
        Inches(7.6), Inches(1.6), Inches(4.9), Inches(4.9),
        "Kishanganga River Corridor",
        "Turquoise waters winding through high mountain valleys in Kashmir."
    )

    # ==================== SLIDE 5: Culture & Dard-Shina People ====================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, COLOR_LIGHT_BG)
    add_header(slide5, "Culture & Heritage", "The Dard-Shina People")

    add_image_card(
        slide5,
        "presentation/images/Local_Culture.jpg",
        Inches(0.8), Inches(1.6), Inches(5.2), Inches(4.9),
        "Dardish Shepherd in Gurez",
        "Local elder wearing traditional wool caps and apparel."
    )

    points_s5 = [
        ("Ancestry & Shina Language", "Inhabited by the Shin tribe, direct descendants of ancient Indo-Aryan Dardic peoples. They speak Shina, a language unique to this border belt."),
        ("Log-Cabin Architecture", "Villages feature multi-story wooden log cabins crafted to withstand heavy winter snowfall and provide natural thermal insulation."),
        ("Cultural Preservation", "Due to decades of isolation, Gurez preserves distinct folk music, traditional dress, and community customs untouched by urban homogenization.")
    ]

    for i, (title, desc) in enumerate(points_s5):
        top_pos = Inches(1.6 + i * 1.7)
        add_card(slide5, Inches(6.3), top_pos, Inches(6.2), Inches(1.5))

        tb = slide5.shapes.add_textbox(Inches(6.5), top_pos + Inches(0.15), Inches(5.8), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_DARK

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_before = Pt(4)

    # ==================== SLIDE 6: Biodiversity & Wildlife ====================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, COLOR_LIGHT_BG)
    add_header(slide6, "Biodiversity", "Flora, Fauna & Wildlife Sanctuary")

    points_s6 = [
        ("Rare Wildlife Sanctuary", "Gurez Valley is a sanctuary for endangered species including the Himalayan Brown Bear, Snow Leopard, Musk Deer, and ibex."),
        ("Sub-Alpine Flora", "Extensive pine, fir, and birch forests intermingled with vibrant wildflower pastures blooming during short summer months."),
        ("Medicinal Herbs", "Rich repository of rare Himalayan herbs such as Black Cumin (Kala Zeera) and valuable medicinal plants used in traditional remedies.")
    ]

    for i, (title, desc) in enumerate(points_s6):
        top_pos = Inches(1.6 + i * 1.7)
        add_card(slide6, Inches(0.8), top_pos, Inches(6.5), Inches(1.5))

        tb = slide6.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.15), Inches(6.1), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_DARK

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_before = Pt(4)

    add_image_card(
        slide6,
        "presentation/images/Gurez_Valley_Scenery.jpg",
        Inches(7.6), Inches(1.6), Inches(4.9), Inches(4.9),
        "Sub-Alpine Meadow Landscape",
        "Coniferous forests and rolling high-altitude pastures."
    )

    # ==================== SLIDE 7: History & Ancient Silk Route ====================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, COLOR_LIGHT_BG)
    add_header(slide7, "History & Heritage", "The Ancient Silk Route Connection")

    add_image_card(
        slide7,
        "presentation/images/Dawar_Town.jpg",
        Inches(0.8), Inches(1.6), Inches(5.2), Inches(4.9),
        "Dawar Town and Bridges",
        "Traditional architecture connecting ancient trade villages."
    )

    points_s7 = [
        ("Ancient Trade Gateway", "Historically part of the Silk Route branch linking Kashmir to Gilgit, Kashgar, and Central Asia."),
        ("Archaeological Treasure", "Inscriptions in ancient Shina and Kharosthi found carved on rocks along the river route trace centuries of trade and pilgrimage."),
        ("Border Realm & Peace Dividend", "Situated near the Line of Control (LoC). Recent years of ceasefire have ushered in an era of peace, tourism, and renewal.")
    ]

    for i, (title, desc) in enumerate(points_s7):
        top_pos = Inches(1.6 + i * 1.7)
        add_card(slide7, Inches(6.3), top_pos, Inches(6.2), Inches(1.5))

        tb = slide7.shapes.add_textbox(Inches(6.5), top_pos + Inches(0.15), Inches(5.8), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_DARK

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_before = Pt(4)

    # ==================== SLIDE 8: Eco Tourism & Adventure ====================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, COLOR_LIGHT_BG)
    add_header(slide8, "Adventure Tourism", "Trekking, Camping & Eco-Exploration")

    points_s8 = [
        ("High-Altitude Trekking", "Base for treks to high alpine tarns such as Satsar Lake, Gadsar Lake, and cross-passes into Drass and Ladakh."),
        ("River Rafting & Camping", "Kishanganga offers thrill seekers white-water rafting opportunities and riverside campsites under starlit skies."),
        ("Offbeat Border Tourism", "Awarded 'Best Offbeat Tourist Destination' in India (2022), offering serene, uncrowded exploration.")
    ]

    for i, (title, desc) in enumerate(points_s8):
        top_pos = Inches(1.6 + i * 1.7)
        add_card(slide8, Inches(0.8), top_pos, Inches(6.5), Inches(1.5))

        tb = slide8.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.15), Inches(6.1), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_DARK

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_before = Pt(4)

    add_image_card(
        slide8,
        "presentation/images/Wooden_Log_Houses.jpg",
        Inches(7.6), Inches(1.6), Inches(4.9), Inches(4.9),
        "Traditional Wooden Architecture",
        "Authentic homestays and timber architecture in Dawar."
    )

    # ==================== SLIDE 9: Modern Developments ====================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, COLOR_LIGHT_BG)
    add_header(slide9, "Future Outlook", "Sustainable Tourism & Infrastructure")

    cards_s9 = [
        ("Eco-Homestays & Local Economy", "Focus on community-driven homestays that empower local Dard-Shina households while providing authentic, sustainable hospitality.", COLOR_PRIMARY),
        ("Year-Round Tunnel Infrastructure", "Proposed Razdan Pass tunnel aims to provide all-weather connectivity, overcoming winter isolation and unlocking economic potential.", COLOR_ACCENT),
        ("Cultural & Heritage Protection", "Strict eco-regulations and cultural protection initiatives ensure tourism growth preserves the delicate mountain ecology.", COLOR_PRIMARY_DARK)
    ]

    for i, (title, desc, accent) in enumerate(cards_s9):
        left_pos = Inches(0.8 + i * 3.95)
        card = add_card(slide9, left_pos, Inches(1.8), Inches(3.7), Inches(4.8))

        # Top color strip
        strip = slide9.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos, Inches(1.8), Inches(3.7), Inches(0.15))
        strip.fill.solid()
        strip.fill.fore_color.rgb = accent
        strip.line.fill.background()

        tb = slide9.shapes.add_textbox(left_pos + Inches(0.2), Inches(2.2), Inches(3.3), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_DARK

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_before = Pt(12)

    # ==================== SLIDE 10: Conclusion ====================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10, COLOR_DARK_BG)

    tb10_head = slide10.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(1.0))
    tf10_h = tb10_head.text_frame
    tf10_h.word_wrap = True

    p = tf10_h.paragraphs[0]
    p.text = "VISITING GUREZ VALLEY"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY

    p2 = tf10_h.add_paragraph()
    p2.text = "Experience the Unspoiled Himalayan Sanctuary"
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(255, 255, 255)

    # Left Info Card
    add_card(slide10, Inches(0.8), Inches(2.0), Inches(5.7), Inches(4.7), bg_color=RGBColor(30, 41, 59), border_color=RGBColor(51, 65, 85))
    tb_l = slide10.shapes.add_textbox(Inches(1.1), Inches(2.2), Inches(5.1), Inches(4.3))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    items_left = [
        ("Best Season to Visit", "June to September — Pleasant climate, blooming alpine pastures, and fully accessible Razdan Pass."),
        ("How to Reach", "123 km from Srinagar (~5-6 hours drive via Bandipora and Razdan Pass). Public buses and private taxis available."),
        ("Permits & Security", "Indian citizens require valid Photo ID; foreign tourists require special permission/permits due to proximity to the border.")
    ]

    for title, desc in items_left:
        p1 = tf_l.add_paragraph() if tf_l.paragraphs[0].text else tf_l.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY
        p1.space_before = Pt(8)

        p2 = tf_l.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = RGBColor(203, 213, 225)
        p2.space_before = Pt(2)

    # Right Summary Card
    add_card(slide10, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.7), bg_color=RGBColor(30, 41, 59), border_color=RGBColor(51, 65, 85))
    tb_r = slide10.shapes.add_textbox(Inches(7.1), Inches(2.2), Inches(5.1), Inches(4.3))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    items_right = [
        ("A Cultural & Natural Gem", "Gurez Valley stands as a testament to Kashmir's natural majesty and the rich, ancient heritage of the Dard-Shina people."),
        ("Preserving Paradise", "As tourism expands, preserving its pristine ecosystem and authentic village culture remains paramount."),
        ("Summary Deck Deliverables", "This presentation suite includes interactive HTML web presentation (`presentation/index.html`), 16:9 PowerPoint (`Gurez_Valley_Presentation.pptx`), and Markdown transcript (`Gurez_Valley_Presentation.md`).")
    ]

    for title, desc in items_right:
        p1 = tf_r.add_paragraph() if tf_r.paragraphs[0].text else tf_r.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_ACCENT
        p1.space_before = Pt(8)

        p2 = tf_r.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = RGBColor(203, 213, 225)
        p2.space_before = Pt(2)

    output_path = "Gurez_Valley_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_presentation()
